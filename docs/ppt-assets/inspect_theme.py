# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu
import os

path = r'E:\xwechat_files\wxid_m2b7myecgtvi22_b160\temp\RWTemp\2026-10\7a394142db6bee2b233450694fd1f030\隐盾安全智能体_正式答辩版.pptx'
p = Presentation(path)

# 主题色
from pptx.oxml.ns import qn
theme = p.slide_masters[0].element.getroottree()
# 直接读theme xml
import zipfile, re
z = zipfile.ZipFile(path)
themes = [n for n in z.namelist() if 'theme' in n and n.endswith('.xml')]
for t in themes[:1]:
    xml = z.read(t).decode('utf-8')
    colors = re.findall(r'<a:(\w+)>.*?val="([0-9A-Fa-f]{6})"', xml)
    print('=== 主题色方案 ===')
    for name, val in colors[:12]:
        print(f'  {name}: #{val}')

print()
# 看前几页形状颜色
for idx in [0, 1, 2, 7]:
    if idx >= len(p.slides): continue
    s = p.slides[idx]
    print(f'=== Slide {idx+1} ===')
    for sh in s.shapes:
        try:
            st = str(sh.shape_type)
            if 'PICTURE' in st:
                print(f'  [图] {sh.name}')
            if sh.has_text_frame and sh.text_frame.text.strip():
                txt = sh.text_frame.text.strip().replace('\n','|')[:50]
                # 取第一个run颜色
                color = None
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        try:
                            if r.font.color and r.font.color.rgb:
                                color = str(r.font.color.rgb)
                        except: pass
                        break
                    if color: break
                print(f'  [文字 #{color}] {txt}')
        except Exception as e:
            pass
    print()
