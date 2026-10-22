# -*- coding: utf-8 -*-
"""生成庆园杯预答辩 5 分钟 PPT（严格按组委会四段式）。

输出：docs/隐盾安全智能体_预答辩.pptx（8 页正文，另附备用页不在本文件）
四段：① 作品简介 → ② 主要功能与预期目标 → ③ 设计方案及技术实现 → ④ 创新及应用前景
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.join("docs", "隐盾安全智能体_预答辩.pptx")

DARK = "0B3D2E"
DARK2 = "0E5C3F"
GREEN = "16A163"
GREEN_D = "0E7A4A"
LIGHT = "F7F8F9"
WHITE = "FFFFFF"
TEXT = "1A1E22"
MUTED = "5A626B"
AMBER = "A9761B"
RED = "C4362F"
LINE = "E1E4E8"
FONT = "Microsoft YaHei"
MONO = "Consolas"

W, H = 13.333, 7.5
M = 0.6

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    return s


def rect(s, x, y, w, h, fill=None, line=None, radius=False, shape=MSO_SHAPE.RECTANGLE):
    if radius and shape == MSO_SHAPE.RECTANGLE:
        shape = MSO_SHAPE.ROUNDED_RECTANGLE
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def txt(s, x, y, w, h, text, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
        font=FONT, line=1.15, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color); r.font.name = font
    return tb


def rich(s, x, y, w, h, items, size=12, color=TEXT, line=1.15):
    """items: [(text, {bold,color,size}), ...] 每项一段。"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, (text, style) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = style.get("line", line)
        p.space_after = Pt(style.get("after", 2))
        r = p.add_run(); r.text = text
        r.font.size = Pt(style.get("size", size))
        r.font.bold = bool(style.get("bold"))
        r.font.color.rgb = RGBColor.from_string(style.get("color", color))
        r.font.name = style.get("font", FONT)
    return tb


def kicker(s, text, color=GREEN_D):
    rect(s, M, 0.42, 0.10, 0.30, fill=color)
    txt(s, M + 0.20, 0.42, 9.5, 0.32, text, size=13, color=color, bold=True)


def title(s, text, y=0.85, size=25, color=TEXT):
    txt(s, M, y, W - 2 * M, 0.9, text, size=size, color=color, bold=True, line=1.12)


def pagenum(s, n):
    txt(s, W - 1.0, H - 0.48, 0.5, 0.3, f"{n:02d}", size=10, color=MUTED,
        align=PP_ALIGN.RIGHT, font=MONO)


def chip(s, x, y, w, text, color=GREEN_D, bg="EDF7F1", h=0.32, size=10.5):
    rect(s, x, y, w, h, fill=bg, line=None, radius=True)
    txt(s, x + 0.10, y + 0.03, w - 0.20, h - 0.04, text, size=size, color=color,
        bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def card(s, x, y, w, h, head, body, accent=GREEN, head_size=14, body_size=11.5):
    rect(s, x, y, w, h, fill=WHITE, line=LINE, radius=True)
    rect(s, x, y, 0.06, h, fill=accent)
    txt(s, x + 0.22, y + 0.14, w - 0.40, 0.32, head, size=head_size, color=accent, bold=True)
    txt(s, x + 0.22, y + 0.52, w - 0.40, h - 0.60, body, size=body_size, color=TEXT, line=1.25)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ───────────────────────── 1 封面 ─────────────────────────
s = slide(DARK)
rect(s, W - 3.6, -1.2, 5.2, 5.2, fill=DARK2, shape=MSO_SHAPE.OVAL)
rect(s, -1.8, H - 2.8, 4.6, 4.6, fill="0A3327", shape=MSO_SHAPE.OVAL)
shield = rect(s, M, 1.35, 0.95, 1.05, fill=GREEN, shape=MSO_SHAPE.PENTAGON)
shield.rotation = 180
txt(s, M + 0.02, 1.55, 0.92, 0.55, "Y", size=26, color=WHITE, bold=True,
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, M, 2.75, 10.5, 0.9, "隐盾安全智能体", size=42, color=WHITE, bold=True)
txt(s, M, 3.75, 10.8, 0.6, "让大模型“看不见”敏感数据，却依然能把活干完",
    size=20, color="CFE8DA")
txt(s, M, 4.55, 11.0, 0.45,
    "YINDUN SECURITY AGENT  ·  PRIVACY-GATEWAY DESKTOP AGENT",
    size=11, color="8FD3B4", font=MONO)
for i, (label, w) in enumerate((("V3.6.19", 1.25), ("本地化部署", 1.55),
                                ("隐私计算 · 桌面智能体", 2.55), ("庆园杯预答辩", 1.85))):
    chip(s, M + sum(x[1] for x in (("", 0), ("", 0))) + i * 2.35, 5.55, w, label,
         color="FFFFFF", bg=DARK2, size=11)
txt(s, M, 6.55, 11.0, 0.4, "2026-10-22  ·  团队与学校（按实际填写）", size=11, color="8FD3B4")
notes(s, "各位老师好，我们是隐盾安全智能体。一句话：让大模型看不见敏感数据，却依然能把合同审查、"
         "文件整理这些活干完。下面按作品简介、主要功能与预期目标、设计方案及技术实现、创新及应用前景四部分汇报。")

# ───────────────────────── 2 作品简介 ─────────────────────────
s = slide()
kicker(s, "一、作品简介")
title(s, "数据越敏感，越需要 AI；却越不能把数据交出去")
leaks = [
    ("出网面", "调用云端模型时，敏感原文离开本机"),
    ("落盘面", "临时文件、缓存、日志里留下副本"),
    ("留痕面", "模型厂商的请求日志与训练回流"),
    ("执行面", "越权写盘、误删、命令注入"),
]
for i, (head, body) in enumerate(leaks):
    x = M + (i % 2) * 3.25
    y = 2.0 + (i // 2) * 1.55
    card(s, x, y, 3.05, 1.35, head, body, accent=RED if i < 2 else AMBER, head_size=13, body_size=10.5)
rect(s, 7.2, 1.95, 5.55, 4.45, fill=LIGHT, line=LINE, radius=True)
txt(s, 7.45, 2.15, 5.05, 0.35, "产品定位", size=12, color=GREEN_D, bold=True)
rich(s, 7.45, 2.55, 5.05, 3.6, [
    ("面向涉密办公场景的", {"size": 15}),
    ("本地化隐私防护桌面智能体", {"size": 15, "bold": True, "color": GREEN_D}),
    ("本机 / 局域网 / 云端双算力；数据进模型前脱敏，回答回本机还原。", {"size": 12, "color": MUTED, "after": 8}),
    ("目标用户", {"size": 12, "bold": True, "color": GREEN_D}),
    ("律所 · 医院 · 制造企业 · 政府外包团队", {"size": 12}),
    ("一句话：不是“少用 AI”，而是“放心用 AI”。", {"size": 12, "bold": True, "color": AMBER}),
])
pagenum(s, 2)
notes(s, "企业想把大模型用在合同、薪酬、病历、客户资料上，但数据只要出了本机就失控：出网、落盘、留痕、执行，"
         "四个环节都是泄露面。隐盾的定位是：面向涉密办公的本地化隐私防护桌面智能体。它不要求企业少用 AI，"
         "而是让企业放心用 AI——数据进模型前脱敏，回答回到本机再还原。主要面向律所、医院、制造企业和政府外包团队。")

# ───────────────────────── 3 功能与目标 ─────────────────────────
s = slide()
kicker(s, "二、主要功能与预期目标")
title(s, "五个功能，四个目标：可用、不可见、可审计、可部署")
funcs = [
    ("脱敏对话", "27 类实体内存内替换为占位符，模型全程只见占位符，回答本机还原"),
    ("文件与办公工具", "15 个工具：列目录/读写/移动/复制/改名/批量改后缀/建目录；默认范围＝程序目录+桌面/文档/下载"),
    ("脱敏知识库", "入库即逐块脱敏，向量库只存占位符；检索结果标注“片段不是全文”"),
    ("安全审计", "HMAC-SHA256 哈希链 + 外部锚点，一键校验、导出"),
    ("企业控制面", "只汇总掩码事件与元数据；策略只能收紧；上报默认关闭"),
]
for i, (head, body) in enumerate(funcs[:3]):
    card(s, M + i * 4.10, 1.95, 3.85, 2.15, head, body, accent=GREEN, head_size=14, body_size=10.5)
for i, (head, body) in enumerate(funcs[3:]):
    card(s, M + i * 4.10, 4.25, 3.85, 1.45, head, body, accent=GREEN_D, head_size=13, body_size=10.5)
goals = [("可用", "本机 4B，无需 GPU"), ("不可见", "已识别实体出网明文 0 例"),
         ("可审计", "每次操作留痕"), ("可部署", "两个免安装 exe")]
for i, (head, body) in enumerate(goals):
    x = M + i * 3.05
    rect(s, x, 5.95, 2.85, 0.85, fill="EDF7F1", line=None, radius=True)
    txt(s, x + 0.12, 6.02, 2.65, 0.28, head, size=11, color=GREEN_D, bold=True)
    txt(s, x + 0.12, 6.30, 2.65, 0.35, body, size=9.5, color=TEXT)
pagenum(s, 3)
notes(s, "功能有五块：第一，脱敏对话，27 类实体在内存里换成占位符；第二，15 个文件与办公工具，"
         "覆盖整理、移动、批量改名；第三，脱敏知识库，入库就脱敏，向量库零明文；第四，安全审计，"
         "哈希链加外部锚点；第五，企业控制面，只上报掩码事件，策略只能收紧。预期目标四个：可用——本机 4B、"
         "不需要 GPU；不可见——对已识别实体，出网明文 0 例；可审计——每次操作留痕；可部署——两个免安装 exe。")

# ───────────────────────── 4 架构 + 网关 ─────────────────────────
s = slide()
kicker(s, "三、设计方案及技术实现")
title(s, "三层解耦 + 内存级脱敏网关 + 四道防线")
layers = [("UI 层", "pywebview / WebView2", DARK2), ("服务层", "AgentService · 工具编排 · 审批", GREEN_D),
          ("安全内核 core/", "脱敏 · 权限 · 审计 · 沙箱（零 Qt）", GREEN)]
for i, (head, body, color) in enumerate(layers):
    y = 2.05 + i * 1.05
    rect(s, M, y, 5.7, 0.82, fill=color, line=None, radius=True)
    txt(s, M + 0.18, y + 0.10, 1.55, 0.30, head, size=12, color=WHITE, bold=True)
    txt(s, M + 1.80, y + 0.12, 3.7, 0.55, body, size=9.5, color="EAF5EF")
    if i < 2:
        txt(s, M + 2.65, y + 0.78, 0.5, 0.25, "↓", size=14, color=MUTED)
rect(s, 6.85, 1.95, 5.9, 2.08, fill=WHITE, line=LINE, radius=True)
txt(s, 7.05, 2.08, 5.5, 0.30, "内存级脱敏网关（三段式）", size=12.5, color=GREEN_D, bold=True)
flow = ["anonymize", "invoke", "deanonymize"]
for i, name in enumerate(flow):
    x = 7.05 + i * 1.85
    rect(s, x, 2.55, 1.65, 0.55, fill="EDF7F1", line=None, radius=True)
    txt(s, x, 2.66, 1.65, 0.35, name, size=10, color=GREEN_D, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    if i < 2:
        txt(s, x + 1.66, 2.62, 0.20, 0.35, "→", size=13, color=MUTED)
txt(s, 7.05, 3.22, 5.5, 0.65,
    "真值只在进程内存；映射表 Fernet 加密，主密钥 DPAPI 保护；占位符带 nonce 防劫持",
    size=9.5, color=MUTED, line=1.2)
txt(s, 6.85, 4.20, 5.9, 0.30, "四道防线", size=12.5, color=GREEN_D, bold=True)
defenses = [("沙箱", "默认可读写范围 + realpath 双重校验"), ("白名单", "去 shell，python 限沙箱内脚本"),
            ("审批", "熔断挂起，展示真实路径与内容"), ("审计", "HMAC 链 + 锚点，可发现截断")]
for i, (head, body) in enumerate(defenses):
    x = 6.85 + (i % 2) * 3.0
    y = 4.55 + (i // 2) * 0.95
    rect(s, x, y, 2.85, 0.82, fill=LIGHT, line=None, radius=True)
    txt(s, x + 0.12, y + 0.08, 0.9, 0.28, head, size=10.5, color=GREEN_D, bold=True)
    txt(s, x + 0.12, y + 0.36, 2.62, 0.42, body, size=9, color=TEXT)
rect(s, M, 6.05, 11.9, 0.70, fill="FDF7EA", line=None, radius=True)
txt(s, M + 0.16, 6.14, 11.6, 0.55,
    "数字口径（带分母，不写 100%）：留出集 60 篇 / 280 个注入值，完整脱敏 264 个、未覆盖 16 个（94.3%），清单公开；"
    "对已识别实体，出网明文 0 例。",
    size=10, color=AMBER, line=1.15)
pagenum(s, 4)
notes(s, "设计上是三层解耦：UI 用 pywebview，服务层做工具编排和审批，安全内核在 core 里、零 Qt、"
         "可以独立测试。核心是内存级脱敏网关：anonymize、invoke、deanonymize 三段式，真值只在进程内存里出现；"
         "映射表 Fernet 加密，主密钥 DPAPI 保护。再往外是四道防线：沙箱、命令白名单、熔断审批、哈希链审计。"
         "数字我们不写 100%：留出集 60 篇、280 个注入值，完整脱敏 264 个，未覆盖 16 个，94.3%；"
         "对已经识别出来的实体，出网明文 0 例。")

# ───────────────────────── 5 laya + 插件 ─────────────────────────
s = slide()
kicker(s, "三、设计方案及技术实现")
title(s, "laya 三级分级 + 插件只做加法：安全能力可扩展，判定权不外包", size=21)
txt(s, M, 1.75, 5.8, 0.30, "laya 风险分级：按数据边界由近及远", size=13, color=GREEN_D, bold=True)
tiers = [("本机 laya 决策模型", "数据不出这台电脑", GREEN),
         ("局域网 laya 服务", "出机器，不出企业内网", "4A7DBF"),
         ("正则规则", "全离线；规则空则如实说“未给出结论”", AMBER)]
for i, (head, body, color) in enumerate(tiers):
    y = 2.15 + i * 1.05
    rect(s, M, y, 5.65, 0.78, fill=LIGHT, line=None, radius=True)
    rect(s, M, y, 0.06, 0.78, fill=color)
    txt(s, M + 0.18, y + 0.08, 5.25, 0.30, head, size=11.5, color=TEXT, bold=True)
    txt(s, M + 0.18, y + 0.40, 5.25, 0.30, body, size=9.5, color=MUTED)
    if i < 2:
        txt(s, M + 2.60, y + 0.78, 0.5, 0.25, "↓ 不可用时自动回落", size=8.5, color=MUTED)
rich(s, M, 5.55, 5.8, 1.1, [
    ("只做建议，不参与放行：", {"bold": True, "color": RED, "size": 10.5}),
    ("高风险只影响弹窗提示，放行仍由人工审批；提示里写明本次来源，绝不冒充模型。",
     {"size": 9.5, "color": TEXT}),
])
txt(s, 6.75, 1.75, 6.0, 0.30, "插件体系：“只做加法”的五条硬约束", size=13, color=GREEN_D, bold=True)
rules = [
    ("7 个内置插件", "敏感词 / 邮件网页解析 / 工作流模板 / 导出格式 / 决策提示 / 操作意图预判 / laya 分级"),
    ("能力只增不减", "多识别、多解析、多导出、多模板；不能让任何内容免于脱敏"),
    ("判定权永久禁止", "approve / mask / bypass / whitelist / write_audit 等钩子不开放"),
    ("工程约束", "载荷净化 · 超时隔离 · 内核兜底 · 启停即收回"),
]
for i, (head, body) in enumerate(rules):
    y = 2.15 + i * 0.92
    txt(s, 6.75, y, 5.95, 0.28, head, size=10.5, color=GREEN_D, bold=True)
    txt(s, 6.75, y + 0.27, 5.95, 0.55, body, size=9, color=TEXT, line=1.15)
rect(s, 6.75, 5.85, 5.95, 0.72, fill=DARK, line=None, radius=True)
txt(s, 6.95, 5.98, 5.6, 0.45, "插件是“能力插座”，不是“权限后门”——能力可以加，安全边界不外包。",
    size=11, color=WHITE, bold=True)
pagenum(s, 5)
notes(s, "这一页讲两个技术实现。左边是 laya 风险分级：按数据边界分三级——本机 laya 模型、局域网服务、"
         "正则规则；每一级都写明本次来源，模型不可用就自动回落，但只做建议、不参与放行。"
         "右边是插件体系：7 个内置插件，可以多识别一类实体、多解析一种格式、多一种导出；但插件永远拿不到判定权——"
         "审批、脱敏豁免、白名单、审计写入，这些钩子在宿主里是永久禁止的。也就是说，能力可以扩展，安全边界不外包。")

# ───────────────────────── 6 四个创新 ─────────────────────────
s = slide()
kicker(s, "四、创新及应用前景")
title(s, "四个创新点：网关、分级、用途、可验证")
inno = [
    ("内存级脱敏网关", "模型物理上看不见真值；机密性是架构保证", "已识别实体出网明文 0 例；500 轮定向劫持未成功（限该攻击面）", GREEN),
    ("laya 三级 + 插件只做加法", "安全能力可按数据边界扩展，判定权永不外包", "7 个插件；判定类钩子永久禁止；启停即收回", "4A7DBF"),
    ("先判用途，再判操作", "先问“想干什么”，再问“对不对得上”，最后问“操作风险”", "个人信息+出售/牟利 → 模型与工具之前直接拒绝；外发 → 红色警示", RED),
    ("可验证交付", "测试通过不等于安全证明，但每一步都能复核", "39 个套件在 CI 通过；verify_archive --all 0 问题；双 exe", AMBER),
]
for i, (head, body, evidence, color) in enumerate(inno):
    x = M + (i % 2) * 6.05
    y = 1.95 + (i // 2) * 2.25
    card(s, x, y, 5.85, 2.05, head, body + "\n\n证据：" + evidence, accent=color,
         head_size=14, body_size=9.5)
pagenum(s, 6)
notes(s, "创新点四个。第一，内存级脱敏网关，机密性是架构保证，不是识别率保证。第二，laya 三级分级加插件只做加法："
         "能力可以扩展，判定权不外包。第三，先判用途，再判操作：以前审批只问“操作对不对得上请求”，"
         "现在先问“这件事该不该做”——收集个人信息去出售，在模型和工具之前就拒绝；外发则红色警示。"
         "第四，可验证交付：39 个套件、双 exe、打包后校验资产等于源码。")

# ───────────────────────── 7 场景/商业/边界 ─────────────────────────
s = slide()
kicker(s, "四、创新及应用前景")
title(s, "场景、商业与边界：能私有化部署，也知道自己不能做什么")
cols = [
    ("应用场景", GREEN, ["律所：合同审查 / 客户资料整理 / 卷宗归档",
                       "医院：病历摘要 / 科研数据整理（本地部署）",
                       "制造企业：供应商合同 / 报价单 / 图纸清单",
                       "政府外包：涉密文档流转前的脱敏与审计"]),
    ("商业与推广", "4A7DBF", ["客户端：私有化部署，按席位/年订阅",
                          "控制面：企业版订阅，集中策略与审计",
                          "行业插件：定制词表 / 解析器 / 工作流模板"]),
    ("边界与路线图", AMBER, ["控制面是演示级原型：RBAC / TLS / 多租户是下一步",
                          "覆盖率是经验值：留出集 94.3%，未覆盖 16 个形态公开",
                          "本机模型能力有限：复杂任务建议切云端或更大模型"]),
]
for i, (head, color, lines) in enumerate(cols):
    x = M + i * 4.10
    rect(s, x, 1.95, 3.85, 3.95, fill=WHITE, line=LINE, radius=True)
    rect(s, x, 1.95, 3.85, 0.06, fill=color)
    txt(s, x + 0.20, 2.12, 3.45, 0.32, head, size=13, color=color, bold=True)
    rich(s, x + 0.20, 2.58, 3.45, 3.15,
         [(ln, {"size": 10, "after": 7}) for ln in lines])
rect(s, M, 6.12, 11.9, 0.62, fill="EDF7F1", line=None, radius=True)
txt(s, M + 0.16, 6.22, 11.6, 0.42,
    "路线图：插件市场与行业模板 → 控制面生产化 → 端侧合规报表", size=11, color=GREEN_D, bold=True)
pagenum(s, 7)
notes(s, "应用上先做律所、医院、制造企业和政府外包，客户端私有化部署，控制面做企业订阅，行业插件做定制。"
         "边界也如实说：控制面还是演示级原型，RBAC、TLS、多租户是下一步；识别覆盖率是经验值，94.3%，"
         "未覆盖的 16 个形态我们公开清单，不承诺 100%。")

# ───────────────────────── 8 收尾 ─────────────────────────
s = slide(DARK)
rect(s, W - 3.2, -1.0, 4.6, 4.6, fill=DARK2, shape=MSO_SHAPE.OVAL)
txt(s, M, 2.15, 11.5, 0.5, "四、创新及应用前景", size=14, color="8FD3B4", bold=True)
txt(s, M, 2.75, 11.5, 1.2, "不是概念，是能跑、能验证、\n能持续交付", size=34, color=WHITE, bold=True, line=1.15)
txt(s, M, 4.55, 11.0, 0.5, "隐盾安全智能体 · V3.6.19", size=14, color="CFE8DA")
txt(s, M, 5.05, 11.0, 0.4, "谢谢各位老师，欢迎提问。", size=16, color=WHITE)
pagenum(s, 8)
notes(s, "我们最想请评委看的是：这套东西不是概念，是能跑、能验证、能持续交付的。谢谢。")

os.makedirs("docs", exist_ok=True)
prs.save(OUT)
print("PPTX saved:", OUT, len(prs.slides.__iter__.__self__._sldIdLst), "slides")
